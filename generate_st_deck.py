from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = "鸿蒙支付开发者效率提升专项_ST汇报模板.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = "172B4D"
BLUE = "1769E0"
CYAN = "00A9CE"
TEAL = "00A884"
GREEN = "27AE60"
ORANGE = "F2994A"
RED = "D64545"
PURPLE = "7B61A8"
INK = "263238"
GRAY = "667085"
MID = "D0D5DD"
LIGHT = "F5F7FA"
PALE_BLUE = "EAF2FF"
PALE_GREEN = "EAF8F3"
PALE_ORANGE = "FFF4E8"
PALE_RED = "FFF0F0"
WHITE = "FFFFFF"
FONT = "Microsoft YaHei"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, font=FONT,
             margin=0.06, line_spacing=1.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.line_spacing = line_spacing
    for run in p.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def add_rich_text(slide, lines, x, y, w, h, size=15, color=INK,
                  bullet=False, spacing=4):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.1)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.08)
    tf.margin_bottom = Inches(0.05)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(line, tuple):
            text, c, is_bold = line
        else:
            text, c, is_bold = line, color, False
        p.text = text
        p.level = 0
        p.space_after = Pt(spacing)
        p.line_spacing = 1.05
        if bullet:
            p.text = "• " + p.text
        for run in p.runs:
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = is_bold
            run.font.color.rgb = rgb(c)
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=MID, radius=True, lw=1):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(lw)
    return shape


def line(slide, x1, y1, x2, y2, color=MID, width=1.5, arrow=False):
    c = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if arrow:
        c.line.end_arrowhead = True
    return c


def title(slide, no, text, subtitle=None):
    add_text(slide, f"{no:02d}", 0.55, 0.35, 0.55, 0.38, 13, BLUE, True,
             PP_ALIGN.CENTER)
    add_text(slide, text, 1.18, 0.25, 11.25, 0.55, 25, NAVY, True)
    line(slide, 0.55, 0.92, 12.75, 0.92, MID, 1)
    if subtitle:
        add_text(slide, subtitle, 1.18, 0.73, 11.0, 0.25, 10, GRAY)


def footer(slide, text="鸿蒙支付开放能力｜开发者效率提升和支撑专项｜内部汇报模板"):
    add_text(slide, text, 0.6, 7.14, 10.8, 0.2, 8.5, GRAY)
    add_text(slide, "所有【待填】内容均可编辑", 10.85, 7.14, 1.9, 0.2,
             8.5, BLUE, False, PP_ALIGN.RIGHT)


def tag(slide, text, x, y, w, fill=PALE_BLUE, color=BLUE):
    rect(slide, x, y, w, 0.34, fill, fill)
    add_text(slide, text, x + 0.03, y + 0.01, w - 0.06, 0.30, 10, color, True,
             PP_ALIGN.CENTER)


def card(slide, x, y, w, h, heading, body, accent=BLUE, fill=WHITE,
         body_size=13):
    rect(slide, x, y, w, h, fill, MID)
    rect(slide, x, y, 0.08, h, accent, accent, radius=False, lw=0)
    add_text(slide, heading, x + 0.23, y + 0.12, w - 0.38, 0.38,
             16, NAVY, True)
    if isinstance(body, list):
        add_rich_text(slide, body, x + 0.2, y + 0.57, w - 0.34, h - 0.67,
                      body_size, GRAY, True)
    else:
        add_text(slide, body, x + 0.22, y + 0.57, w - 0.38, h - 0.69,
                 body_size, GRAY, False, PP_ALIGN.LEFT, MSO_ANCHOR.TOP)


def metric_card(slide, x, y, w, label, value="【待填】", note="口径：【待填】",
                accent=BLUE):
    rect(slide, x, y, w, 1.15, WHITE, MID)
    add_text(slide, label, x + 0.16, y + 0.10, w - 0.32, 0.28, 11, GRAY)
    add_text(slide, value, x + 0.16, y + 0.37, w - 0.32, 0.40, 23, accent, True)
    add_text(slide, note, x + 0.16, y + 0.82, w - 0.32, 0.22, 8.5, GRAY)


def add_notes(slide, text):
    # A visible, removable speaker-note strip keeps the file portable and editable.
    rect(slide, 9.6, 6.58, 3.15, 0.42, "FFFBEA", "F0D35E")
    add_text(slide, "讲述提示：" + text, 9.72, 6.62, 2.92, 0.32, 8.5, GRAY)


# 1. Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, NAVY)
rect(slide, 0, 0, 0.18, 7.5, CYAN, CYAN, radius=False, lw=0)
tag(slide, "研发 ST 会议｜可编辑模板", 0.72, 0.65, 2.25, fill="25466F", color=WHITE)
add_text(slide, "鸿蒙支付开发者效率提升\n与支撑专项", 0.72, 1.45, 8.4, 1.7,
         34, WHITE, True, valign=MSO_ANCHOR.TOP)
add_text(slide, "从“人工经验驱动”升级为标准化、工具化、智能化、可度量的支付接入交付体系",
         0.75, 3.35, 8.7, 0.7, 18, "D9E8FF")
rect(slide, 9.65, 1.25, 2.75, 3.75, "203A62", "315681")
for idx, (a, b) in enumerate([
    ("商户接入", "更快"),
    ("研发交付", "更稳"),
    ("技术支撑", "可规模化"),
]):
    yy = 1.65 + idx * 1.02
    add_text(slide, a, 9.95, yy, 1.1, 0.30, 11, "B8CBE5")
    add_text(slide, b, 9.95, yy + 0.28, 1.95, 0.45, 21, WHITE, True)
add_text(slide, "汇报人：【待填】    部门：【待填】    日期：【待填】",
         0.75, 6.55, 8.2, 0.35, 12, "B8CBE5")

# 2. Executive summary
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
title(slide, 2, "专项摘要：这次要解决什么问题")
rect(slide, 0.62, 1.18, 1.22, 0.42, BLUE, BLUE)
add_text(slide, "一句话目标", 0.65, 1.22, 1.15, 0.32, 11, WHITE, True,
         PP_ALIGN.CENTER)
rect(slide, 1.95, 1.18, 10.75, 0.92, PALE_BLUE, "B9D4FF")
add_text(slide,
         "将支付接入从依赖人工经验和一对一支撑，升级为标准化、工具化、智能化、可度量的开发者交付体系。",
         2.22, 1.31, 10.1, 0.58, 20, NAVY, True)
for i, (h, b, c) in enumerate([
    ("为什么做", ["接入链路长、文档与能力分散", "签名验签及回调问题高频", "支撑经验难复用、成本持续增长"], RED),
    ("重点做什么", ["AI Skill 与场景化接入", "SDK / Sample / 文档一体化", "联调、上线门禁与诊断能力"], BLUE),
    ("交付价值", ["缩短首笔成功支付耗时", "提升首次成功率与上线质量", "提高自助解决率、降低工单"], TEAL),
]):
    card(slide, 0.65 + i * 4.12, 2.48, 3.82, 2.55, h, b, c, WHITE, 13)
rect(slide, 0.65, 5.38, 12.0, 0.92, "F8FAFC", MID)
add_text(slide, "本次 ST 希望达成", 0.88, 5.54, 1.55, 0.30, 12, NAVY, True)
add_text(slide, "① 目标与指标口径确认   ② 五大专项优先级确认   ③ 跨团队责任与资源确认   ④ 环境与数据能力决策",
         2.55, 5.47, 9.55, 0.45, 14, INK)
add_notes(slide, "开门见山讲清楚目标、价值和本次会议要决策的事项。")
footer(slide)

# 3. Journey pain points
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
title(slide, 3, "现状诊断：商户接入旅程与关键断点",
      "建议补充真实数据、典型商户案例和 Top 问题分布")
steps = [
    ("① 接入准备", "入网 / AppID\n证书 / 权限", RED),
    ("② 方案设计", "商户模型\n场景选型", ORANGE),
    ("③ 编码集成", "端侧 / 服务端\n回调处理", BLUE),
    ("④ 联调验证", "异常场景\n生产小额联调", PURPLE),
    ("⑤ 上线验收", "安全 / 幂等\n压测 / 参数", TEAL),
    ("⑥ 生产运营", "监控 / 对账\n退款 / 排障", GREEN),
]
for i, (h, b, c) in enumerate(steps):
    x = 0.55 + i * 2.08
    rect(slide, x, 1.35, 1.78, 1.18, WHITE, c)
    add_text(slide, h, x + 0.10, 1.48, 1.58, 0.28, 12, c, True,
             PP_ALIGN.CENTER)
    add_text(slide, b, x + 0.10, 1.80, 1.58, 0.55, 11, GRAY, False,
             PP_ALIGN.CENTER)
    if i < len(steps) - 1:
        line(slide, x + 1.78, 1.94, x + 2.07, 1.94, MID, 2, True)
pain_cards = [
    ("配置链路复杂", "主体、商户号、AppID、证书、产品权限涉及多平台协同。"),
    ("知识资产分散", "客户端、服务端、回调与商户模型文档缺少一站式路径。"),
    ("安全实现门槛高", "签名、SM2 验签、私钥保护、幂等处理依赖开发经验。"),
    ("联调保障不足", "Payment Kit 暂无明确免扣费独立沙盒，生产联调风险较高。"),
    ("问题定位效率低", "错误码缺少上下文，重复咨询与人工排障比例高。"),
    ("过程不可度量", "尚未形成从接入资格到首笔成功支付的统一漏斗。"),
]
for i, (h, b) in enumerate(pain_cards):
    row, col = divmod(i, 3)
    card(slide, 0.65 + col * 4.1, 3.05 + row * 1.35, 3.78, 1.08,
         h, b, [RED, ORANGE, PURPLE][col], "FAFBFC", 10.5)
add_notes(slide, "用数据回答：问题发生在哪一步、影响多少商户、造成多少支撑成本。")
footer(slide)

# 4. Goals and metrics
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
title(slide, 4, "目标体系：以首笔成功支付耗时为北极星指标",
      "指标必须同时给出基线、目标、数据源和责任人")
rect(slide, 0.65, 1.22, 12.0, 1.18, NAVY, NAVY)
add_text(slide, "北极星指标", 0.93, 1.42, 1.25, 0.28, 11, "BFD6F4", True)
add_text(slide, "Time to First Successful Payment", 2.22, 1.31, 4.65, 0.45,
         23, WHITE, True)
add_text(slide, "商户获得接入资格 → 完成首笔服务端确认的有效支付",
         6.95, 1.42, 5.25, 0.34, 13, "D9E8FF")
metrics = [
    ("效率", "平均接入周期", BLUE),
    ("质量", "首次回调验签成功率", TEAL),
    ("支撑", "开发者自助解决率", PURPLE),
    ("稳定性", "订单状态一致率", GREEN),
]
for i, (group, label, c) in enumerate(metrics):
    x = 0.65 + i * 3.08
    tag(slide, group, x, 2.78, 0.78, fill=c, color=WHITE)
    metric_card(slide, x, 3.20, 2.83, label, "基线【待填】",
                "目标：【待填】｜数据源：【待填】", c)
    metric_card(slide, x, 4.58, 2.83,
                ["首笔支付耗时", "上线检查通过率", "单商户工单量", "生产问题率"][i],
                "目标【待填】", "责任人：【待填】", c)
rect(slide, 0.65, 6.10, 12.0, 0.55, PALE_ORANGE, "FFD4A6")
add_text(slide, "若当前没有可靠基线：第一阶段先完成旅程埋点、指标定义和数据看板，再承诺改善幅度。",
         0.88, 6.20, 11.55, 0.30, 12, "805100", True)
add_notes(slide, "不要只报工具数量；用效率、质量、支撑成本和稳定性衡量结果。")
footer(slide)

# 5. Panorama
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
title(slide, 5, "能力全景：覆盖研发、联调、上线与运营闭环")
stage_names = ["接入准备", "方案设计", "编码集成", "联调验证", "上线验收", "生产运营"]
for i, name in enumerate(stage_names):
    x = 0.55 + i * 2.08
    rect(slide, x, 1.15, 1.78, 0.52, BLUE if i < 5 else TEAL,
         BLUE if i < 5 else TEAL)
    add_text(slide, f"{i+1:02d}  {name}", x + 0.05, 1.22, 1.68, 0.32,
             11, WHITE, True, PP_ALIGN.CENTER)
rows = [
    ("研发辅助工具", ["接入指南\n配置清单", "Payment Skill\n场景任务", "@kit.PaymentKit\npay-java / REST", "错误码诊断\nSample / Mock", "上线检查清单\n参数复核", "FAQ / 工单\n知识回流"], PALE_BLUE, BLUE),
    ("安全与质量", ["身份与权限\n检查", "商户模型\n合规选型", "签名 / 验签\n敏感信息加密", "异常场景\n回调幂等测试", "私钥 / HTTPS\n压测验收", "订单一致性\n生产监控"], PALE_RED, RED),
    ("平台与支撑", ["商户平台\nAppGallery Connect", "文档 / API\n方案咨询", "客户端 / 服务端\n示例工程", "生产小额联调\n专家支持", "准入门禁\n上线护航", "报表 / 退款\n对账 / 排障"], PALE_GREEN, TEAL),
]
for r, (label, cells, fill, c) in enumerate(rows):
    y = 1.95 + r * 1.36
    rect(slide, 0.55, y, 1.22, 1.08, c, c)
    add_text(slide, label, 0.63, y + 0.13, 1.06, 0.78, 12, WHITE, True,
             PP_ALIGN.CENTER)
    for i, content in enumerate(cells):
        x = 1.92 + i * 1.78
        rect(slide, x, y, 1.57, 1.08, fill, "D9E1EA")
        add_text(slide, content, x + 0.06, y + 0.10, 1.45, 0.86, 10.5, INK,
                 False, PP_ALIGN.CENTER)
rect(slide, 0.55, 6.15, 12.0, 0.52, "F8FAFC", MID)
add_text(slide, "闭环机制：接入数据与工单 → Top 问题识别 → 文档 / SDK / Skill / 工具改进 → 版本发布 → 效果度量",
         0.82, 6.25, 11.5, 0.30, 12, NAVY, True, PP_ALIGN.CENTER)
add_notes(slide, "强调不是单点 SDK 建设，而是覆盖开发者全生命周期的交付体系。")
footer(slide)

# 6. Five workstreams
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
title(slide, 6, "五大关键专项：从单点工具走向体系化交付")
projects = [
    ("01", "AI 接入提效", "Payment Skill 2.0\n场景任务编排\n代码生成与检查", BLUE),
    ("02", "SDK 易用性", "Java SDK 增强\n语言覆盖评估\n统一异常与安全封装", CYAN),
    ("03", "联调工具", "签名验证\n回调模拟\n异常场景测试", PURPLE),
    ("04", "上线保障", "自动检查清单\n准入门禁\n首发护航与监控", RED),
    ("05", "支撑数字化", "开发者旅程埋点\n接入看板\n工单与知识闭环", TEAL),
]
for i, (no, h, b, c) in enumerate(projects):
    x = 0.55 + i * 2.52
    rect(slide, x, 1.35, 2.22, 4.35, WHITE, MID)
    rect(slide, x, 1.35, 2.22, 0.74, c, c)
    add_text(slide, no, x + 0.16, 1.48, 0.42, 0.34, 17, WHITE, True)
    add_text(slide, h, x + 0.62, 1.48, 1.42, 0.34, 14, WHITE, True)
    add_text(slide, b, x + 0.20, 2.37, 1.82, 1.35, 13, INK, False,
             PP_ALIGN.CENTER)
    line(slide, x + 0.25, 3.95, x + 1.97, 3.95, MID, 1)
    add_text(slide, "阶段交付", x + 0.25, 4.10, 1.70, 0.25, 10, GRAY, True)
    add_text(slide, "【待填】\n\n验收指标：【待填】\n责任团队：【待填】",
             x + 0.25, 4.37, 1.70, 1.05, 10.5, c)
rect(slide, 0.65, 6.02, 12.0, 0.58, PALE_BLUE, "B9D4FF")
add_text(slide, "优先级建议：先做可度量的基础设施与高频痛点，再扩展长尾场景和更多语言 SDK。",
         0.92, 6.13, 11.45, 0.32, 12, NAVY, True)
add_notes(slide, "逐项讲清楚问题、交付物、验收指标和责任团队。")
footer(slide)

# 7. AI and developer assets
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
title(slide, 7, "重点专项一：AI 接入与开发资产一体化")
card(slide, 0.65, 1.25, 3.78, 4.85, "Payment Skill 2.0",
     ["按商户模型生成接入任务清单", "覆盖支付、签约、退款等场景", "生成端侧与服务端代码骨架",
      "内置安全红线与错误诊断", "生成结果可追溯、可审查"], BLUE, PALE_BLUE, 13)
card(slide, 4.78, 1.25, 3.78, 4.85, "SDK / Sample 工程",
     ["@kit.PaymentKit 易用性优化", "pay-java 业务与异常模型增强", "主流技术栈最小可运行 Demo",
      "签名、回调、幂等参考实现", "评估 Node.js / Go 等语言需求"], CYAN, "ECFAFD", 13)
card(slide, 8.90, 1.25, 3.78, 4.85, "文档 / 诊断资产",
     ["按开发者旅程重构文档导航", "场景化快速接入与可复制示例", "错误码—原因—检查—修复映射",
      "配置检查与回调诊断工具", "文档、SDK、Skill 版本一致性"], TEAL, PALE_GREEN, 13)
add_text(slide, "验收结果", 0.75, 6.33, 0.88, 0.25, 10, BLUE, True)
add_text(slide, "接入步骤减少【待填】｜代码复用率【待填】｜首次成功率提升【待填】｜自助解决率【待填】",
         1.70, 6.25, 10.65, 0.42, 13, NAVY, True)
add_notes(slide, "Skill 是开发辅助，不替代商户审核、安全验收和上线门禁。")
footer(slide)

# 8. Testing and release assurance
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
title(slide, 8, "重点专项二：联调、上线门禁与生产兜底")
columns = [
    ("联调验证", ["本地签名 / 验签工具", "回调 Mock 与重复回调", "成功、失败、取消、超时场景", "并发、乱序与重试测试", "生产小额受控联调"], PURPLE),
    ("上线门禁", ["商户号与 AppID 绑定", "证书、密钥及产品权限", "HTTPS 回调连通性", "先验签后更新订单", "幂等、压测与参数复核"], RED),
    ("生产兜底", ["回调与主动查询互补", "订单状态一致性检测", "商户维度异常监控", "首发护航与升级通道", "退款、对账及问题复盘"], TEAL),
]
for i, (h, items, c) in enumerate(columns):
    x = 0.65 + i * 4.12
    card(slide, x, 1.32, 3.78, 3.85, h, items, c, WHITE, 13)
    rect(slide, x + 0.82, 5.29, 2.12, 0.58, c, c)
    add_text(slide, ["开发完成", "允许上线", "持续稳定"][i],
             x + 0.95, 5.35, 1.85, 0.45, 15, WHITE, True, PP_ALIGN.CENTER)
    if i < 2:
        line(slide, x + 3.20, 5.58, x + 4.0, 5.58, MID, 2, True)
rect(slide, 0.65, 6.18, 12.0, 0.50, PALE_RED, "FFC7C7")
add_text(slide,
         "现状边界：Payment Kit 暂无明确免真实扣费独立沙盒；需同步推进受控测试环境或等价模拟能力评估。",
         0.90, 6.28, 11.55, 0.28, 11.5, "8A2525", True)
add_notes(slide, "把上线保障讲成可执行门禁，而不是仅依赖文档提醒。")
footer(slide)

# 9. Support operating model
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
title(slide, 9, "支撑模式升级：从人肉答疑到分层、自助和知识闭环")
levels = [
    ("L0", "自助知识", "文档 / FAQ / Sample / 错误码", BLUE),
    ("L1", "智能诊断", "Payment Skill / 配置检查 / 问题定位", CYAN),
    ("L2", "标准支持", "结构化工单 / 一线技术支持", TEAL),
    ("L3", "研发专家", "复杂接入 / 疑难问题 / 专项治理", ORANGE),
    ("L4", "联合保障", "重大生产问题 / 跨团队应急", RED),
]
for i, (lv, h, b, c) in enumerate(levels):
    y = 1.25 + i * 0.94
    rect(slide, 0.65 + i * 0.18, y, 6.0 - i * 0.36, 0.70, c, c)
    add_text(slide, lv, 0.82 + i * 0.18, y + 0.13, 0.48, 0.32, 14, WHITE, True)
    add_text(slide, h, 1.35 + i * 0.18, y + 0.13, 1.10, 0.32, 13, WHITE, True)
    add_text(slide, b, 2.62 + i * 0.18, y + 0.13, 3.45 - i * 0.36, 0.32,
             11, WHITE)
rect(slide, 7.15, 1.25, 5.5, 4.45, "F8FAFC", MID)
add_text(slide, "持续改进闭环", 7.48, 1.50, 2.0, 0.35, 18, NAVY, True)
loop = [("接入数据与工单", BLUE), ("Top 问题识别", PURPLE),
        ("工具与知识改进", ORANGE), ("发布与效果度量", TEAL)]
for i, (txt, c) in enumerate(loop):
    angle_pos = [(7.55, 2.25), (10.10, 2.25), (10.10, 4.05), (7.55, 4.05)][i]
    rect(slide, angle_pos[0], angle_pos[1], 2.12, 0.78, WHITE, c)
    add_text(slide, txt, angle_pos[0] + 0.12, angle_pos[1] + 0.18, 1.88, 0.34,
             12, c, True, PP_ALIGN.CENTER)
line(slide, 9.67, 2.64, 10.08, 2.64, MID, 2, True)
line(slide, 11.16, 3.03, 11.16, 4.03, MID, 2, True)
line(slide, 10.08, 4.44, 9.69, 4.44, MID, 2, True)
line(slide, 8.61, 4.03, 8.61, 3.05, MID, 2, True)
rect(slide, 0.65, 6.15, 12.0, 0.52, PALE_GREEN, "B9EAD9")
add_text(slide, "目标：【待填】自助解决率｜【待填】工单下降率｜【待填】平均响应时长｜Top 问题闭环周期【待填】",
         0.88, 6.25, 11.55, 0.30, 12, NAVY, True, PP_ALIGN.CENTER)
add_notes(slide, "高频问题必须回灌产品资产，避免研发专家长期重复答疑。")
footer(slide)

# 10. Roadmap
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
title(slide, 10, "交付路线：按能力阶段推进并设置可验收出口")
phases = [
    ("阶段 A｜基线与高频问题", ["统一指标口径与旅程埋点", "梳理 Top 问题与关键断点", "完善基础支付 Skill 与检查清单"], BLUE),
    ("阶段 B｜工具与门禁", ["联调、签名与回调诊断工具", "SDK / Sample / 文档版本协同", "上线门禁及首发护航机制"], PURPLE),
    ("阶段 C｜规模化与闭环", ["扩展支付场景及技术栈", "支撑分层与知识自动回流", "数据驱动专项持续治理"], TEAL),
]
for i, (h, items, c) in enumerate(phases):
    x = 0.65 + i * 4.12
    card(slide, x, 1.28, 3.78, 2.35, h, items, c, WHITE, 12)
    rect(slide, x, 3.86, 3.78, 1.65, "F8FAFC", MID)
    add_text(slide, "阶段出口", x + 0.20, 4.03, 1.0, 0.28, 11, c, True)
    add_text(slide, "交付物：【待填】\n验收指标：【待填】\nOwner：【待填】",
             x + 0.20, 4.35, 3.25, 0.88, 11, INK)
    if i < 2:
        line(slide, x + 3.78, 2.48, x + 4.10, 2.48, MID, 2, True)
rect(slide, 0.65, 5.88, 12.0, 0.72, PALE_ORANGE, "FFD4A6")
add_text(slide, "里程碑填写原则", 0.88, 6.07, 1.28, 0.30, 11, "805100", True)
add_text(slide, "使用“能力完成 + 指标达到”的验收出口；具体日期、版本和团队排期由项目组补充。",
         2.27, 6.00, 9.95, 0.40, 12, "805100")
add_notes(slide, "避免只列发布时间；每阶段都应有可验证的开发者结果。")
footer(slide)

# 11. Decisions and risks
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
title(slide, 11, "风险、依赖与需要 ST 决策的事项")
card(slide, 0.65, 1.25, 3.78, 4.92, "主要风险",
     ["缺少统一数据导致价值难量化", "多平台、多团队依赖影响交付闭环", "AI 生成结果存在安全与版本风险",
      "测试环境不足影响联调效率", "专项变成工具堆砌而非结果改进"], RED, PALE_RED, 13)
card(slide, 4.78, 1.25, 3.78, 4.92, "关键依赖",
     ["商户平台与 AGC 配置数据", "支付链路日志与埋点", "工单及知识库数据",
      "安全、测试、运营团队协同", "SDK、文档、Skill 统一发布机制"], ORANGE, PALE_ORANGE, 13)
card(slide, 8.90, 1.25, 3.78, 4.92, "建议 ST 决策",
     ["确认北极星指标及统一口径", "确认五大专项优先级与 Owner", "决策受控沙盒 / 模拟环境方向",
      "决策服务端语言 SDK 投入", "确认上线门禁与跨团队资源"], BLUE, PALE_BLUE, 13)
rect(slide, 0.65, 6.38, 12.0, 0.38, NAVY, NAVY)
add_text(slide, "会议结论记录：【待填】",
         0.88, 6.41, 11.2, 0.25, 11.5, WHITE, True)
add_notes(slide, "收口到明确决策、责任人和下一步，而不是只做信息同步。")
footer(slide)

# 12. Appendix / fill guide
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
title(slide, 12, "附录：模板填写清单")
items = [
    ("业务现状", "商户数量、典型接入周期、Top 失败原因、工单量及案例", BLUE),
    ("目标指标", "基线、目标、统计口径、数据源、Owner", TEAL),
    ("专项交付", "范围、非范围、关键能力、阶段出口、责任团队", PURPLE),
    ("上线保障", "门禁项、验收方式、例外机制、生产护航方案", RED),
    ("资源与决策", "跨团队依赖、人员与环境投入、需要 ST 决策事项", ORANGE),
]
for i, (h, b, c) in enumerate(items):
    y = 1.20 + i * 0.98
    rect(slide, 0.70, y, 0.65, 0.65, c, c)
    add_text(slide, f"{i+1:02d}", 0.78, y + 0.12, 0.48, 0.30, 13, WHITE, True,
             PP_ALIGN.CENTER)
    add_text(slide, h, 1.60, y + 0.05, 1.35, 0.30, 14, NAVY, True)
    add_text(slide, b, 3.10, y + 0.03, 8.85, 0.45, 12, GRAY)
    line(slide, 1.58, y + 0.73, 12.35, y + 0.73, "E5EAF0", 1)
rect(slide, 0.70, 6.35, 11.95, 0.46, "F8FAFC", MID)
add_text(slide, "编辑说明：可直接替换【待填】文字；所有图形、文本框、流程线及指标卡均为 PPT 原生可编辑元素。",
         0.92, 6.43, 11.50, 0.28, 11, BLUE, True)
footer(slide)

prs.save(OUT)
print(f"Generated {OUT} with {len(prs.slides)} slides")
